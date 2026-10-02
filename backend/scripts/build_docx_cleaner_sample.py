"""Regenerate the downloadable Docx Cleaner sample.

    cd backend && .venv/Scripts/python scripts/build_docx_cleaner_sample.py

Has all three things the tool can strip: document properties (author etc.), a real comment
(python-docx 1.2+ has native support), and a tracked insertion/deletion (python-docx has no API
for this, so it is built with raw OXML - the same technique `docx_cleaner.docx_service` undoes).
"""

import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent.parent / "app" / "tools" / "docx_cleaner" / "samples"


def _tracked_run(tag: str, text: str, author: str, when: datetime.datetime) -> OxmlElement:
    """<w:ins>/<w:del> wrapping one run; w:del uses <w:delText> instead of <w:t>."""
    wrapper = OxmlElement(f"w:{tag}")
    wrapper.set(qn("w:id"), "1")
    wrapper.set(qn("w:author"), author)
    wrapper.set(qn("w:date"), when.strftime("%Y-%m-%dT%H:%M:%SZ"))
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:delText" if tag == "del" else "w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    wrapper.append(run)
    return wrapper


def build() -> Document:
    doc = Document()
    doc.core_properties.author = "Budi Santoso"
    doc.core_properties.last_modified_by = "Budi Santoso"
    doc.core_properties.comments = "Internal draft - do not distribute"
    doc.core_properties.keywords = "draft, internal, confidential"
    doc.core_properties.title = "Docx Cleaner sample"

    doc.add_heading("Docx Cleaner Sample", level=1)
    doc.add_paragraph(
        "This document has hidden metadata, a reviewer comment and a tracked change - try "
        "cleaning it and compare the result in Word's File > Info panel and Review tab."
    )

    p = doc.add_paragraph("The budget for this project is ")
    now = datetime.datetime(2026, 1, 15, 9, 30, 0)
    p._p.append(_tracked_run("ins", "finalized and approved", "Sari Wulandari", now))
    p._p.append(_tracked_run("del", "still being negotiated", "Sari Wulandari", now))
    p.add_run(".")

    commented = doc.add_paragraph()
    r1 = commented.add_run("This paragraph has a reviewer ")
    r2 = commented.add_run("comment attached")
    commented.add_run(" to it.")
    doc.add_comment(
        [r1, r2],
        text="Please double-check this figure before sending externally.",
        author="Andi Pratama",
        initials="AP",
    )

    return doc


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build().save(OUT / "sample.docx")
    print(f"Wrote {OUT / 'sample.docx'}")


if __name__ == "__main__":
    main()
