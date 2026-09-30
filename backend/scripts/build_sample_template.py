"""Regenerate the downloadable Text Replacer samples.

    cd backend && .venv/Scripts/python scripts/build_sample_template.py

Writes app/tools/text_replacer/samples/{template.docx,data.csv}. The CSV header must list
exactly the variables used in the template (tests/test_samples.py enforces this).
"""

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "app" / "tools" / "text_replacer" / "samples"
FONT = "Times New Roman"
ACCENT = RGBColor(0x1F, 0x38, 0x64)


def set_font(run, size=None, bold=None, underline=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if underline is not None:
        run.underline = underline
    if color is not None:
        run.font.color.rgb = color


def add(par, text, **fmt):
    """Add text as ONE run. Keep each {{placeholder}} inside a single run like this."""
    run = par.add_run(text)
    set_font(run, **fmt)
    return run


def bottom_border(par, color="1F3864", size="12"):
    pbdr = OxmlElement("w:pBdr")
    line = OxmlElement("w:bottom")
    for key, value in {"w:val": "single", "w:sz": size, "w:space": "4", "w:color": color}.items():
        line.set(qn(key), value)
    pbdr.append(line)
    par._p.get_or_add_pPr().append(pbdr)


def spacing(par, before=0, after=6, line=1.15):
    fmt = par.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(before), Pt(after), line


def build_template() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(12)

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)

    # Letterhead lives in the header: placeholders in headers/footers are replaced too.
    head = section.header.paragraphs[0]
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add(head, "{{perusahaan}}", size=18, bold=True, color=ACCENT)
    spacing(head, after=0)
    sub = section.header.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add(sub, "{{alamat_perusahaan}}", size=9, color=RGBColor(0x55, 0x55, 0x55))
    bottom_border(sub)
    spacing(sub, after=0)

    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add(foot, "Nomor: {{nomor_surat}}", size=9, color=RGBColor(0x77, 0x77, 0x77))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add(title, "SURAT KETERANGAN KERJA", size=14, bold=True, underline=True)
    spacing(title, before=12, after=0)
    number = doc.add_paragraph()
    number.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add(number, "Nomor: {{nomor_surat}}")
    spacing(number, after=18)

    intro = doc.add_paragraph()
    add(intro, "Yang bertanda tangan di bawah ini menerangkan bahwa:")
    spacing(intro, after=8)

    # Placeholders in table cells work like body text.
    fields = [
        ("Nama", "{{nama}}"),
        ("NIK Karyawan", "{{nik_karyawan}}"),
        ("Jabatan", "{{jabatan}}"),
        ("Departemen", "{{departemen}}"),
        ("Tanggal Bergabung", "{{tanggal_bergabung}}"),
    ]
    table = doc.add_table(rows=len(fields), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Cm(4.5), Cm(0.6), Cm(10.4))
    table.autofit = False
    for column, width in zip(table.columns, widths, strict=True):
        column.width = width  # also updates the grid, which LibreOffice/Word honour
    for row, (label, placeholder) in zip(table.rows, fields, strict=True):
        for cell, width, text in zip(row.cells, widths, (label, ":", placeholder), strict=True):
            cell.width = width
            par = cell.paragraphs[0]
            add(par, text, bold=(text == placeholder and label == "Nama"))
            spacing(par, after=3)

    body = doc.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add(body, "Adalah benar karyawan ")
    add(body, "{{perusahaan}}", bold=True)
    add(body, " yang telah bekerja sejak tanggal {{tanggal_bergabung}} hingga saat ini ")
    add(body, "pada posisi {{jabatan}}.")
    spacing(body, before=14, after=8)

    purpose = doc.add_paragraph()
    purpose.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add(purpose, "Surat keterangan ini dibuat atas permintaan yang bersangkutan untuk keperluan ")
    add(purpose, "{{keperluan}}", bold=True)  # bold here = value comes out bold
    add(purpose, " dan dapat dipergunakan sebagaimana mestinya.")
    spacing(purpose, after=28)

    for text, kwargs in [
        ("{{kota}}, {{tanggal_surat}}", {}),
        ("Hormat kami,", {}),
    ]:
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(9)
        add(par, text, **kwargs)
        spacing(par, after=0)
    gap = doc.add_paragraph()
    gap.paragraph_format.left_indent = Cm(9)
    spacing(gap, before=40, after=0)
    signer = doc.add_paragraph()
    signer.paragraph_format.left_indent = Cm(9)
    add(signer, "{{penandatangan}}", bold=True, underline=True)
    spacing(signer, after=0)
    role = doc.add_paragraph()
    role.paragraph_format.left_indent = Cm(9)
    add(role, "{{jabatan_penandatangan}}")
    spacing(role, after=0)

    doc.core_properties.title = "Contoh template Text Replacer"
    doc.core_properties.author = "Shorui"
    return doc


SAMPLE_ROWS = [
    {
        "perusahaan": "PT Contoh Sejahtera",
        "alamat_perusahaan": "Jl. Merdeka No. 10, Jakarta Pusat 10110 | Telp. (021) 555-0100",
        "nomor_surat": "001/HRD/SKK/IX/2026",
        "nama": "Budi Santoso",
        "nik_karyawan": "EMP-0142",
        "jabatan": "Senior Software Engineer",
        "departemen": "Teknologi Informasi",
        "tanggal_bergabung": "1 Maret 2021",
        "keperluan": "pengajuan kredit kepemilikan rumah",
        "kota": "Jakarta",
        "tanggal_surat": "30 September 2026",
        "penandatangan": "Siti Rahayu",
        "jabatan_penandatangan": "Manajer HRD",
    },
    {
        "nomor_surat": "002/HRD/SKK/IX/2026",
        "nama": "Dewi Lestari",
        "nik_karyawan": "EMP-0207",
        "jabatan": "Product Designer",
        "departemen": "Desain Produk",
        "tanggal_bergabung": "15 Agustus 2022",
        "keperluan": "pengajuan visa kunjungan",
    },
    {
        "nomor_surat": "003/HRD/SKK/IX/2026",
        "nama": "Andi Wijaya",
        "nik_karyawan": "EMP-0318",
        "jabatan": "Quality Assurance Analyst",
        "departemen": "Quality Assurance",
        "tanggal_bergabung": "2 Januari 2024",
        "keperluan": "pembukaan rekening bank",
    },
]


def build_csv(columns: list[str]) -> None:
    defaults = SAMPLE_ROWS[0]
    shared = ["perusahaan", "alamat_perusahaan", "kota", "tanggal_surat", "penandatangan",
              "jabatan_penandatangan"]  # fmt: skip
    with (OUT / "data.csv").open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns)
        writer.writeheader()
        for row in SAMPLE_ROWS:
            writer.writerow({c: row.get(c, defaults[c] if c in shared else "") for c in columns})


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_template().save(OUT / "template.docx")
    # Import lazily so the script also works before the app is importable elsewhere.
    from app.tools.text_replacer.docx_service import extract_variables

    columns = extract_variables((OUT / "template.docx").read_bytes())
    build_csv(columns)
    print(f"Wrote {OUT} ({len(columns)} variables: {', '.join(columns)})")


if __name__ == "__main__":
    main()
