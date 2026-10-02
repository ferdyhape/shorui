"""Regenerate the downloadable Docx to PDF sample.

    cd backend && .venv/Scripts/python scripts/build_docx_to_pdf_sample.py

Plain prose (no {{placeholders}}) so converting it demonstrates layout, not mail-merge.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = Path(__file__).resolve().parent.parent / "app" / "tools" / "docx_to_pdf" / "samples"


def build() -> Document:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Quarterly Project Update")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Prepared for the Shorui documentation tools demo").italic = True

    doc.add_paragraph(
        "This is a plain Word document with no merge fields - it exists to show what Docx to "
        "PDF does: render a .docx exactly as Word would print it (fonts, spacing, page breaks, "
        "tables) into a PDF nobody can accidentally edit."
    )

    doc.add_heading("Status by workstream", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Workstream", "Status", "Owner"
    for workstream, status, owner in [
        ("Backend API", "On track", "Budi"),
        ("Frontend UI", "On track", "Sari"),
        ("Documentation", "Needs review", "Andi"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = workstream, status, owner

    doc.add_heading("Next steps", level=2)
    for item in ("Finish the remaining review comments", "Ship the beta build", "Collect feedback"):
        doc.add_paragraph(item, style="List Bullet")

    doc.core_properties.title = "Quarterly Project Update (sample)"
    doc.core_properties.author = "Shorui"
    return doc


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build().save(OUT / "sample.docx")
    print(f"Wrote {OUT / 'sample.docx'}")


if __name__ == "__main__":
    main()
