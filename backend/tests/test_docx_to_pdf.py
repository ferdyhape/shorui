import shutil
import subprocess
import tempfile
from pathlib import Path

import pytest
from docx import Document

from app.core.config import Settings
from app.core.errors import ConversionError, ServiceUnavailableError
from app.tools.docx_to_pdf import converter


def make_docx() -> bytes:
    import io

    doc = Document()
    doc.add_paragraph("Hello, PDF.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_resolve_soffice_rejects_bad_configured_path():
    with pytest.raises(ServiceUnavailableError):
        converter.resolve_soffice(Settings(soffice_path="Z:/nope/soffice.exe"))


def test_resolve_soffice_accepts_configured_path(tmp_path: Path):
    fake = tmp_path / "soffice.exe"
    fake.write_bytes(b"")
    assert converter.resolve_soffice(Settings(soffice_path=str(fake))) == str(fake)


def test_convert_raises_service_unavailable_when_soffice_missing(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(converter.shutil, "which", lambda _: None)
    monkeypatch.setattr(converter, "_CANDIDATES", ())
    with pytest.raises(ServiceUnavailableError):
        converter.convert_to_pdf(make_docx(), Settings(), tmp_path)


def test_convert_raises_conversion_error_on_nonzero_exit(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(converter, "resolve_soffice", lambda _: "soffice")
    monkeypatch.setattr(
        converter.subprocess,
        "run",
        lambda *a, **k: subprocess.CompletedProcess([], 1, b"", b"boom"),
    )
    with pytest.raises(ConversionError, match="boom"):
        converter.convert_to_pdf(make_docx(), Settings(), tmp_path)


def test_convert_raises_conversion_error_on_timeout(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(converter, "resolve_soffice", lambda _: "soffice")

    def _raise(*_a, **_k):
        raise subprocess.TimeoutExpired(cmd="soffice", timeout=1)

    monkeypatch.setattr(converter.subprocess, "run", _raise)
    with pytest.raises(ConversionError, match="timed out"):
        converter.convert_to_pdf(make_docx(), Settings(), tmp_path)


def test_convert_raises_service_unavailable_when_exec_not_found(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(converter, "resolve_soffice", lambda _: "soffice")

    def _raise(*_a, **_k):
        raise FileNotFoundError("no such file")

    monkeypatch.setattr(converter.subprocess, "run", _raise)
    with pytest.raises(ServiceUnavailableError):
        converter.convert_to_pdf(make_docx(), Settings(), tmp_path)


def _soffice_available() -> bool:
    if shutil.which("soffice") or shutil.which("libreoffice"):
        return True
    return any(Path(c).is_file() for c in converter._CANDIDATES)


_HAS_SOFFICE = _soffice_available()


@pytest.mark.skipif(not _HAS_SOFFICE, reason="LibreOffice not installed on this machine")
def test_real_conversion_produces_a_pdf():
    with tempfile.TemporaryDirectory() as tmp:
        pdf_bytes = converter.convert_to_pdf(make_docx(), Settings(), Path(tmp))
    assert pdf_bytes.startswith(b"%PDF")


@pytest.mark.skipif(not _HAS_SOFFICE, reason="LibreOffice not installed on this machine")
def test_api_convert_endpoint(client):
    res = client.post(
        "/api/v1/docx-to-pdf/convert",
        files={"file": ("letter.docx", make_docx())},
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")
    assert "letter.pdf" in res.headers["content-disposition"]


def test_api_rejects_non_docx(client):
    res = client.post(
        "/api/v1/docx-to-pdf/convert",
        files={"file": ("letter.txt", b"hi")},
    )
    assert (res.status_code, res.json()["code"]) == (400, "invalid_file")


def test_sample_docx_exists_and_is_plain_prose():
    from docx import Document

    from app.tools.docx_to_pdf.router import SAMPLE_PATH

    doc = Document(SAMPLE_PATH)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{" not in text  # not a Text Replacer template - plain prose for a conversion demo
    assert "Quarterly Project Update" in text


def test_download_sample(client):
    res = client.get("/api/v1/docx-to-pdf/samples/sample.docx")
    assert res.status_code == 200
    assert "sample.docx" in res.headers["content-disposition"]


@pytest.mark.skipif(not _HAS_SOFFICE, reason="LibreOffice not installed on this machine")
def test_real_conversion_of_the_sample_produces_a_pdf():
    from app.tools.docx_to_pdf.router import SAMPLE_PATH

    with tempfile.TemporaryDirectory() as tmp:
        pdf_bytes = converter.convert_to_pdf(SAMPLE_PATH.read_bytes(), Settings(), Path(tmp))
    assert pdf_bytes.startswith(b"%PDF")
