import datetime
import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.tools.docx_cleaner import docx_service
from app.tools.docx_cleaner.docx_service import CleanOptions

API = "/api/v1/docx-cleaner"


def tracked_run(tag: str, text: str) -> OxmlElement:
    wrapper = OxmlElement(f"w:{tag}")
    wrapper.set(qn("w:id"), "1")
    wrapper.set(qn("w:author"), "Reviewer")
    wrapper.set(qn("w:date"), "2026-01-01T00:00:00Z")
    run = OxmlElement("w:r")
    el = OxmlElement("w:delText" if tag == "del" else "w:t")
    el.set(qn("xml:space"), "preserve")
    el.text = text
    run.append(el)
    wrapper.append(run)
    return wrapper


def make_doc_with_everything() -> bytes:
    doc = Document()
    doc.core_properties.author = "Budi Santoso"
    doc.core_properties.last_modified_by = "Budi Santoso"
    doc.core_properties.keywords = "confidential"
    doc.core_properties.created = datetime.datetime(2020, 1, 1)  # left alone by design

    p = doc.add_paragraph("Status: ")
    p._p.append(tracked_run("ins", "approved"))
    p._p.append(tracked_run("del", "pending"))

    commented = doc.add_paragraph()
    r1 = commented.add_run("flagged text")
    doc.add_comment([r1], text="check this", author="Andi")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_strip_properties_only():
    data = make_doc_with_everything()
    out = docx_service.clean(
        data, CleanOptions(strip_properties=True, strip_comments=False, accept_revisions=False)
    )
    doc = Document(io.BytesIO(out))
    assert doc.core_properties.author == ""
    assert doc.core_properties.last_modified_by == ""
    assert doc.core_properties.keywords == ""
    assert doc.core_properties.created.replace(tzinfo=None) == datetime.datetime(2020, 1, 1)
    assert len(list(doc.comments)) == 1  # untouched


def test_accept_revisions_keeps_insertions_drops_deletions():
    data = make_doc_with_everything()
    out = docx_service.clean(
        data, CleanOptions(strip_properties=False, strip_comments=False, accept_revisions=True)
    )
    text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
    assert "Status: approved" in text
    assert "pending" not in text
    # ins/del wrapper elements themselves must be gone, not just invisible
    doc = Document(io.BytesIO(out))
    xml = doc.element.body.xml
    assert "w:ins" not in xml
    assert "w:del" not in xml


def test_strip_comments_removes_anchors_but_keeps_text():
    data = make_doc_with_everything()
    out = docx_service.clean(
        data, CleanOptions(strip_properties=False, strip_comments=True, accept_revisions=False)
    )
    doc = Document(io.BytesIO(out))
    assert "flagged text" in "\n".join(p.text for p in doc.paragraphs)
    xml = doc.element.body.xml
    assert "commentReference" not in xml
    assert "commentRangeStart" not in xml


def test_all_options_together():
    data = make_doc_with_everything()
    out = docx_service.clean(data, CleanOptions())
    doc = Document(io.BytesIO(out))
    assert doc.core_properties.author == ""
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "approved" in text and "pending" not in text


def test_api_clean_endpoint(client):
    data = make_doc_with_everything()
    res = client.post(f"{API}/clean", files={"file": ("secret.docx", data)})
    assert res.status_code == 200
    assert "secret-cleaned.docx" in res.headers["content-disposition"]
    doc = Document(io.BytesIO(res.content))
    assert doc.core_properties.author == ""


def test_api_rejects_no_options_selected(client):
    data = make_doc_with_everything()
    res = client.post(
        f"{API}/clean",
        files={"file": ("a.docx", data)},
        data={"strip_properties": "false", "strip_comments": "false", "accept_revisions": "false"},
    )
    assert (res.status_code, res.json()["code"]) == (400, "invalid_rows")


def test_download_sample(client):
    res = client.get(f"{API}/samples/sample.docx")
    assert res.status_code == 200
    assert "sample.docx" in res.headers["content-disposition"]


def test_download_sample_rejects_unknown_name(client):
    assert client.get(f"{API}/samples/secret.docx").status_code == 422


def test_sample_has_properties_comments_and_revisions():
    from app.tools.docx_cleaner.router import SAMPLES_DIR

    data = (SAMPLES_DIR / "sample.docx").read_bytes()
    doc = Document(io.BytesIO(data))
    assert doc.core_properties.author
    assert len(list(doc.comments)) >= 1
    assert "w:ins" in doc.element.body.xml
    assert "w:del" in doc.element.body.xml
