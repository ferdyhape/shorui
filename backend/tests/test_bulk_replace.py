import io
import json
import zipfile

from docx import Document

from app.tools.bulk_replace import docx_service
from tests.conftest import all_text


def make_doc(text_runs: list[str], header: str | None = None) -> bytes:
    doc = Document()
    p = doc.add_paragraph()
    for run_text in text_runs:
        p.add_run(run_text)
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_replace_across_split_runs():
    # "Acme Corp" split over two runs, like Word's spellcheck often does.
    data = make_doc(["Dear ", "Ac", "me Corp", ", your order shipped."])
    out = docx_service.replace(data, *docx_service.build_pattern([("Acme Corp", "Globex Inc")]))
    assert "Dear Globex Inc, your order shipped." in all_text(out)


def test_longest_match_wins_over_substring():
    data = make_doc(["Order Total: 100, Order: A1"])
    out = docx_service.replace(
        data,
        *docx_service.build_pattern([("Order", "ORD"), ("Order Total", "GRAND TOTAL")]),
    )
    assert "GRAND TOTAL: 100, ORD: A1" in all_text(out)


def test_replace_in_header_too():
    data = make_doc(["body text"], header="Confidential - Acme")
    out = docx_service.replace(data, *docx_service.build_pattern([("Acme", "Globex")]))
    doc = Document(io.BytesIO(out))
    assert doc.sections[0].header.paragraphs[0].text == "Confidential - Globex"


def test_illegal_xml_characters_stripped():
    data = make_doc(["x"])
    out = docx_service.replace(data, *docx_service.build_pattern([("x", "a\x00b")]))
    assert "ab" in all_text(out)


def test_process_single_file_returns_docx():
    data = make_doc(["hello world"])
    result = docx_service.process([("note.docx", data)], [("hello", "hi")])
    assert result.filename == "note.docx"
    assert "hi world" in all_text(result.content.read())


def test_process_multi_file_returns_zip_and_dedupes_names():
    a = make_doc(["foo"])
    b = make_doc(["foo bar"])
    result = docx_service.process([("note.docx", a), ("note.docx", b)], [("foo", "baz")])
    with zipfile.ZipFile(io.BytesIO(result.content.read())) as zf:
        names = zf.namelist()
        assert names == ["note.docx", "note-2.docx"]
        assert "baz" in all_text(zf.read("note.docx"))
        assert "baz bar" in all_text(zf.read("note-2.docx"))


def test_api_process_single_file(client):
    data = make_doc(["Dear Customer, welcome."])
    res = client.post(
        "/api/v1/bulk-replace/process",
        files=[("files", ("letter.docx", data))],
        data={"pairs": json.dumps([["Customer", "Budi"]])},
    )
    assert res.status_code == 200
    assert "letter.docx" in res.headers["content-disposition"]
    assert "Dear Budi, welcome." in all_text(res.content)


def test_api_rejects_empty_find(client):
    res = client.post(
        "/api/v1/bulk-replace/process",
        files=[("files", ("a.docx", make_doc(["x"])))],
        data={"pairs": json.dumps([["", "y"]])},
    )
    assert (res.status_code, res.json()["code"]) == (400, "invalid_rows")


def test_api_rejects_too_many_files(client):
    # settings fixture caps max_files at 3
    files = [("files", (f"{i}.docx", make_doc(["x"]))) for i in range(4)]
    res = client.post(
        "/api/v1/bulk-replace/process",
        files=files,
        data={"pairs": json.dumps([["x", "y"]])},
    )
    assert (res.status_code, res.json()["code"]) == (400, "invalid_file")


def test_sample_letters_contain_acme_corp_in_header_and_split_runs():
    from app.tools.bulk_replace.router import SAMPLES_DIR

    for name in ("letter-budi.docx", "letter-sari.docx"):
        data = (SAMPLES_DIR / name).read_bytes()
        doc = Document(io.BytesIO(data))
        assert doc.sections[0].header.paragraphs[0].text == "Acme Corp - Customer Service"
        assert "Acme Corp" in all_text(data)  # survives the split-run paragraph too


def test_sample_letters_replace_cleanly(client):
    from app.tools.bulk_replace.router import SAMPLES_DIR

    files = [
        ("files", (name, (SAMPLES_DIR / name).read_bytes()))
        for name in ("letter-budi.docx", "letter-sari.docx")
    ]
    res = client.post(
        "/api/v1/bulk-replace/process",
        files=files,
        data={"pairs": json.dumps([["Acme Corp", "Globex Inc"]])},
    )
    assert res.status_code == 200
    with zipfile.ZipFile(io.BytesIO(res.content)) as zf:
        for name in zf.namelist():
            text = all_text(zf.read(name))
            assert "Globex Inc" in text
            assert "Acme Corp" not in text


def test_download_sample_letters(client):
    for name in ("letter-budi.docx", "letter-sari.docx"):
        res = client.get(f"/api/v1/bulk-replace/samples/{name}")
        assert res.status_code == 200
        assert name in res.headers["content-disposition"]


def test_download_sample_rejects_unknown_name(client):
    assert client.get("/api/v1/bulk-replace/samples/secret.docx").status_code == 422
