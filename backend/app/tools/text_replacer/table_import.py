"""Parse CSV / XLSX uploads into (columns, rows-of-strings)."""

import csv
from collections.abc import Iterable, Iterator
from contextlib import contextmanager
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.config import Settings
from app.core.errors import InvalidFileError, InvalidTableError


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise InvalidTableError("Cannot decode CSV file")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@contextmanager
def _csv_rows(data: bytes) -> Iterator[Iterable[Iterable[Any]]]:
    text = _decode(data)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    yield csv.reader(StringIO(text, newline=""), dialect)


@contextmanager
def _xlsx_rows(data: bytes) -> Iterator[Iterable[Iterable[Any]]]:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise InvalidFileError("File is not a valid .xlsx workbook") from exc
    try:
        if wb.active is None:
            raise InvalidTableError("Workbook has no sheets")
        yield wb.active.iter_rows(values_only=True)
    finally:
        wb.close()  # read-only mode keeps the zip handle open otherwise


def parse_table(
    filename: str, data: bytes, settings: Settings
) -> tuple[list[str], list[dict[str, str]]]:
    lower = filename.lower()
    if lower.endswith((".csv", ".txt")):
        opener = _csv_rows
    elif lower.endswith((".xlsx", ".xlsm")):
        opener = _xlsx_rows
    else:
        raise InvalidTableError("Unsupported file type. Use .csv or .xlsx")

    try:
        with opener(data) as raw_rows:
            return _read(iter(raw_rows), settings)
    except csv.Error as exc:
        raise InvalidTableError(f"Malformed CSV: {exc}") from exc


def _read(
    iterator: Iterator[Iterable[Any]], settings: Settings
) -> tuple[list[str], list[dict[str, str]]]:
    header_raw = next(iterator, None)
    if header_raw is None:
        raise InvalidTableError("File is empty")

    columns: list[str] = []
    positions: list[tuple[int, str]] = []
    for i, cell in enumerate(header_raw):
        name = _cell(cell)
        if name and name not in columns:
            columns.append(name)
            positions.append((i, name))
    if not columns:
        raise InvalidTableError("First row must contain column headers")
    if len(columns) > settings.max_columns:
        raise InvalidTableError(f"Too many columns (max {settings.max_columns})")

    rows: list[dict[str, str]] = []
    for raw in iterator:
        cells = list(raw)
        row = {name: _cell(cells[i]) if i < len(cells) else "" for i, name in positions}
        if not any(row.values()):
            continue
        rows.append(row)
        if len(rows) > settings.max_rows:  # stop parsing early on huge files
            raise InvalidTableError(f"Too many rows (max {settings.max_rows})")
    return columns, rows
