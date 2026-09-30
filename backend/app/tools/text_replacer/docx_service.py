"""Find and replace {{variable}} placeholders in .docx files.

Word often splits one placeholder over several runs (<w:t> nodes), e.g. after
spellcheck or partial formatting. So matching runs on the joined paragraph
text, then writes the value back into the first touched node. Formatting of the
first node wins.
"""

import re
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from io import BytesIO
from tempfile import SpooledTemporaryFile
from typing import IO, Any

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from app.core.errors import InvalidFileError

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ZIP_MIME = "application/zip"

VAR_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_W_P = qn("w:p")
_W_T = qn("w:t")
_INVALID_FILENAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
# Characters XML 1.0 cannot represent; lxml raises ValueError if they reach a node.
_ILLEGAL_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")
_ZIP_SPOOL_BYTES = 8 * 1024 * 1024
TOOL_SLUG = "text-replacer"


def _load(data: bytes) -> Any:
    try:
        return Document(BytesIO(data))
    except (
        PackageNotFoundError,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        AttributeError,
        etree.XMLSyntaxError,
    ) as exc:
        raise InvalidFileError("File is not a valid .docx document") from exc


def _roots(doc: Any) -> Iterator[Any]:
    """Body plus every distinct header/footer part."""
    yield doc.element.body
    seen: set[int] = set()
    for section in doc.sections:
        parts = (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        )
        for part in parts:
            if part.is_linked_to_previous:
                continue
            element = part._element  # python-docx exposes no public accessor
            if id(element) in seen:
                continue
            seen.add(id(element))
            yield element


def _paragraphs(doc: Any) -> Iterator[Any]:
    # Iterating w:p covers body, nested tables and text boxes.
    for root in _roots(doc):
        yield from root.iter(_W_P)


def _text_nodes(paragraph: Any) -> list[Any]:
    """w:t nodes owned by this paragraph (not by a nested text-box paragraph)."""
    nodes = []
    for t in paragraph.iter(_W_T):
        owner = t.getparent()
        while owner is not None and owner.tag != _W_P:
            owner = owner.getparent()
        if owner is paragraph:
            nodes.append(t)
    return nodes


def _write_text(node: Any, text: str) -> None:
    """Set node text; newlines become <w:br/> so Word shows real line breaks."""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    node.text = lines[0]
    node.set(_XML_SPACE, "preserve")
    anchor = node
    for line in lines[1:]:
        br = OxmlElement("w:br")
        t = OxmlElement("w:t")
        t.text = line
        t.set(_XML_SPACE, "preserve")
        anchor.addnext(br)
        br.addnext(t)
        anchor = t


def _replace_in_paragraph(paragraph: Any, values: dict[str, str]) -> None:
    nodes = _text_nodes(paragraph)
    if not nodes:
        return
    original = [n.text or "" for n in nodes]
    matches = list(VAR_RE.finditer("".join(original)))
    if not matches:
        return

    texts = list(original)
    lengths = [len(t) for t in original]
    starts = []
    pos = 0
    for length in lengths:
        starts.append(pos)
        pos += length

    # Reverse order keeps earlier offsets valid while text changes.
    for m in reversed(matches):
        s, e = m.span()
        value = _ILLEGAL_XML.sub("", values.get(m.group(1), ""))
        hit = [i for i in range(len(nodes)) if starts[i] + lengths[i] > s and starts[i] < e]
        if not hit:
            continue
        first, last = hit[0], hit[-1]
        prefix = texts[first][: s - starts[first]]
        suffix = texts[last][e - starts[last] :]
        if first == last:
            texts[first] = prefix + value + suffix
        else:
            texts[first] = prefix + value
            for i in hit[1:-1]:
                texts[i] = ""
            texts[last] = suffix

    for node, old, new in zip(nodes, original, texts, strict=True):
        if new != old:
            _write_text(node, new)


def extract_variables(data: bytes) -> list[str]:
    """Unique variable names, in order of first appearance."""
    doc = _load(data)
    found: dict[str, None] = {}
    for paragraph in _paragraphs(doc):
        text = "".join(t.text or "" for t in _text_nodes(paragraph))
        for m in VAR_RE.finditer(text):
            found.setdefault(m.group(1), None)
    return list(found)


def render(data: bytes, values: dict[str, str]) -> bytes:
    """Fill placeholders. Variables missing from `values` become empty."""
    doc = _load(data)
    for paragraph in _paragraphs(doc):
        _replace_in_paragraph(paragraph, values)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _safe_name(text: str) -> str:
    return _INVALID_FILENAME.sub("_", text).strip(" .")[:80]


def template_stem(template_name: str) -> str:
    return _safe_name(re.sub(r"\.docx$", "", template_name, flags=re.IGNORECASE)) or "document"


def zip_name(template_name: str) -> str:
    """Archive name: <tool>_<template name>.zip, e.g. text-replacer_offer-letter.zip."""
    return f"{TOOL_SLUG}_{template_stem(template_name)}.zip"


def output_names(
    template_name: str, rows: list[dict[str, str]], filename_key: str | None
) -> list[str]:
    """One unique file name per row: <key value>_<template name>.docx (or <name>_<n>)."""
    stem = template_stem(template_name)
    used: set[str] = set()
    names = []
    for index, row in enumerate(rows, start=1):
        if filename_key:
            prefix = _safe_name(row.get(filename_key, "").strip()) or f"row{index}"
            base = f"{prefix}_{stem}"
        else:
            base = f"{stem}_{index}"
        name, n = f"{base}.docx", 2
        while name.lower() in used:
            name = f"{base}-{n}.docx"
            n += 1
        used.add(name.lower())
        names.append(name)
    return names


@dataclass
class GeneratedFile:
    filename: str
    media_type: str
    content: IO[bytes]  # positioned at 0; caller closes
    size: int


def generate(
    template_name: str,
    data: bytes,
    rows: list[dict[str, str]],
    filename_key: str | None,
) -> GeneratedFile:
    """Render one document per row. Several rows are zipped into a spooled temp file
    so memory stays bounded however many documents are produced."""
    names = output_names(template_name, rows, filename_key)
    if len(rows) == 1:
        content = render(data, rows[0])
        return GeneratedFile(names[0], DOCX_MIME, BytesIO(content), len(content))

    spool = SpooledTemporaryFile(max_size=_ZIP_SPOOL_BYTES)  # noqa: SIM115 - closed by caller
    try:
        with zipfile.ZipFile(spool, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, row in zip(names, rows, strict=True):
                zf.writestr(name, render(data, row))
        size = spool.tell()
        spool.seek(0)
    except BaseException:
        spool.close()
        raise
    return GeneratedFile(zip_name(template_name), ZIP_MIME, spool, size)
